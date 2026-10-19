"""
Excel Loader for US DOE TECHTEST 1.0 workbooks.
Parses filled .xlsx templates into FullEntryInput or RapidEntryInput models.
"""

from pathlib import Path
from typing import Union, List, Optional
import openpyxl
from pytechtest.core.types import (
    TRLLevel,
    CapExItem,
    CapExCategory,
    MaterialItem,
    EnergyItem,
    InUseItem,
    GridScenario
)
from pytechtest.core.opex import LaborInput
from pytechtest.models.project import ProjectMetadata, MarketDefinition, FunctionalUnitDefinition
from pytechtest.models.full_entry_input import FullEntryInput, SubsystemConfig
from pytechtest.models.rapid_entry_input import RapidEntryInput, RapidTechCard


class TechTestExcelLoader:
    """Parses official US DOE TECHTEST 1.0 Excel files."""

    def __init__(self, file_path: Union[str, Path]):
        self.file_path = Path(file_path)
        if not self.file_path.exists():
            raise FileNotFoundError(f"TECHTEST Excel file not found: {self.file_path}")
        self.wb = openpyxl.load_workbook(str(self.file_path), data_only=True)

    def load_project_metadata(self) -> ProjectMetadata:
        """Parse Sheet 4 'Project Basics & Benchmarks'."""
        ws = self.wb["Project Basics & Benchmarks"]
        title = str(ws["C7"].value or "Untitled TECHTEST Project").strip()
        lead_org = str(ws["C8"].value or "").strip()
        partners = str(ws["C9"].value or "").strip()
        pi = str(ws["C10"].value or "").strip()
        pi_org = str(ws["C11"].value or "").strip()
        amo_mgr = str(ws["C12"].value or "").strip()
        start = str(ws["C13"].value or "").strip()
        end = str(ws["C14"].value or "").strip()
        funding = str(ws["C15"].value or "").strip()

        # Parse TRL if present
        trl_val = TRLLevel.TRL_3
        for r in range(16, 22):
            val = str(ws.cell(r, 3).value or "")
            if "TRL" in val:
                try:
                    num = int(val.split()[1].replace(":", ""))
                    trl_val = TRLLevel(num)
                    break
                except (IndexError, ValueError):
                    pass

        return ProjectMetadata(
            title=title,
            lead_organization=lead_org,
            partners=partners,
            principal_investigator=pi,
            pi_organization=pi_org,
            technology_manager=amo_mgr,
            start_date=start,
            end_date=end,
            funding_mechanism=funding,
            trl=trl_val
        )

    def load_market(self) -> MarketDefinition:
        """Parse Sheet 5 'Market'."""
        ws = self.wb["Market"]
        us_vol = float(ws["C8"].value or 1000.0)
        unit = str(ws["D8"].value or "Million Gallons").strip()
        product = str(ws["E8"].value or ws["C8"].value or "Generic Industrial Product").strip()
        # Clean product string if formatted as ' of ... per year'
        if product.startswith("of "):
            product = product[3:].strip()
        if product.endswith("per year"):
            product = product[:-8].strip()
        
        pen = float(ws["F20"].value or 0.10) if ws["F20"].value is not None else 0.10
        desc = str(ws["C12"].value or "").strip()

        return MarketDefinition(
            generic_end_use_product=product,
            annual_us_market_volume=us_vol,
            market_units=unit,
            target_penetration_fraction=pen,
            market_description=desc
        )

    def load_functional_unit(self) -> FunctionalUnitDefinition:
        """Parse Sheet 6 'Functional Unit'."""
        ws = self.wb["Functional Unit"]
        ref_vol = float(ws["D20"].value or 1.0)
        unit = str(ws["E20"].value or "Million Gallons").strip()
        p_new = str(ws["G20"].value or (ws["D9"].value if "D9" in ws else "New Technology Product")).strip()
        if p_new.startswith("of "):
            p_new = p_new[3:].strip()

        ref_vol_bm = float(ws["D29"].value or ref_vol)
        p_bench = str(ws["G29"].value or (ws["L9"].value if "L9" in ws else "Benchmark Product")).strip()
        if p_bench.startswith("of "):
            p_bench = p_bench[3:].strip()

        equiv_ratio = ref_vol_bm / ref_vol if ref_vol > 0 else 1.0

        return FunctionalUnitDefinition(
            product_name_new_tech=p_new,
            product_name_benchmark=p_bench,
            reference_volume=ref_vol,
            reference_volume_unit=unit,
            equivalence_ratio=equiv_ratio
        )

    def load_full_input(self) -> FullEntryInput:
        """Parse all sheets to construct FullEntryInput."""
        proj = self.load_project_metadata()
        mkt = self.load_market()
        fu = self.load_functional_unit()

        # --- CapEx ---
        ws_cap = self.wb["CapEx"]
        fac_prod_nt = float(ws_cap["C15"].value or 1.0)
        fac_prod_bm = float(ws_cap["K15"].value or ws_cap["C15"].value or 1.0)

        capex_nt: List[CapExItem] = []
        capex_bm: List[CapExItem] = []
        for r in range(22, 55):
            # New Tech
            cat_nt = ws_cap.cell(r, 2).value or CapExCategory.PROCESS_EQUIPMENT
            name_nt = ws_cap.cell(r, 4).value
            price_nt = ws_cap.cell(r, 5).value
            life_nt = ws_cap.cell(r, 6).value
            if name_nt and price_nt is not None and life_nt:
                try:
                    capex_nt.append(CapExItem(
                        name=str(name_nt).strip(),
                        category=CapExCategory.PROCESS_EQUIPMENT,
                        purchase_price=float(price_nt),
                        lifetime_years=float(life_nt)
                    ))
                except (ValueError, TypeError):
                    pass

            # Benchmark
            name_bm = ws_cap.cell(r, 12).value
            price_bm = ws_cap.cell(r, 13).value
            life_bm = ws_cap.cell(r, 14).value
            if name_bm and price_bm is not None and life_bm:
                try:
                    capex_bm.append(CapExItem(
                        name=str(name_bm).strip(),
                        category=CapExCategory.PROCESS_EQUIPMENT,
                        purchase_price=float(price_bm),
                        lifetime_years=float(life_bm)
                    ))
                except (ValueError, TypeError):
                    pass

        # --- OpEx Labor & Recurring ---
        ws_op = self.wb["OpEx"]
        fte_nt = float(ws_op["C22"].value or 0.0)
        salary_nt = float(ws_op["D22"].value or 75000.0)
        fringe_nt = float(ws_op["E22"].value or 0.72)
        labor_nt = LaborInput(fte_count=fte_nt, average_annual_salary=salary_nt, fringe_overhead_rate=fringe_nt)

        fte_bm = float(ws_op["J22"].value or 0.0)
        salary_bm = float(ws_op["K22"].value or 75000.0)
        fringe_bm = float(ws_op["L22"].value or 0.72)
        labor_bm = LaborInput(fte_count=fte_bm, average_annual_salary=salary_bm, fringe_overhead_rate=fringe_bm)

        other_rec_nt = float(ws_op["F69"].value or 0.0) if ws_op["F69"].value is not None else 0.0
        other_rec_bm = float(ws_op["M69"].value or 0.0) if ws_op["M69"].value is not None else 0.0

        # --- Raw Materials ---
        ws_mat = self.wb["Raw Materials"]
        mat_nt: List[MaterialItem] = []
        mat_bm: List[MaterialItem] = []
        for r in range(18, 45):
            # New Tech
            m_name = ws_mat.cell(r, 2).value
            m_qty = ws_mat.cell(r, 3).value
            m_unit = ws_mat.cell(r, 4).value or "kg"
            m_ee = ws_mat.cell(r, 5).value or 0.0
            m_gwp = ws_mat.cell(r, 6).value or 0.0
            m_p = ws_mat.cell(r, 7).value or 0.0
            if m_name and m_qty is not None:
                try:
                    mat_nt.append(MaterialItem(
                        name=str(m_name).strip(),
                        amount_per_ref_volume=float(m_qty),
                        unit=str(m_unit).strip(),
                        unit_price=float(m_p) if m_p else 0.0,
                        embodied_energy_factor=float(m_ee) if m_ee else 0.0,
                        embodied_ghg_factor=float(m_gwp) if m_gwp else 0.0
                    ))
                except (ValueError, TypeError):
                    pass

            # Benchmark
            bm_name = ws_mat.cell(r, 12).value
            bm_qty = ws_mat.cell(r, 13).value
            bm_unit = ws_mat.cell(r, 14).value or "kg"
            bm_ee = ws_mat.cell(r, 15).value or 0.0
            bm_gwp = ws_mat.cell(r, 16).value or 0.0
            bm_p = ws_mat.cell(r, 17).value or 0.0
            if bm_name and bm_qty is not None:
                try:
                    mat_bm.append(MaterialItem(
                        name=str(bm_name).strip(),
                        amount_per_ref_volume=float(bm_qty),
                        unit=str(bm_unit).strip(),
                        unit_price=float(bm_p) if bm_p else 0.0,
                        embodied_energy_factor=float(bm_ee) if bm_ee else 0.0,
                        embodied_ghg_factor=float(bm_gwp) if bm_gwp else 0.0
                    ))
                except (ValueError, TypeError):
                    pass

        # --- Manufacturing Energy ---
        ws_en = self.wb["Manufacturing Energy"]
        en_nt: List[EnergyItem] = []
        en_bm: List[EnergyItem] = []
        for r in range(17, 30):
            src_nt = ws_en.cell(r, 3).value
            amt_nt = ws_en.cell(r, 4).value
            p_nt = ws_en.cell(r, 5).value
            gwp_nt = ws_en.cell(r, 6).value
            if src_nt and amt_nt is not None:
                try:
                    en_nt.append(EnergyItem(
                        energy_source=str(src_nt).strip(),
                        amount_per_ref_volume=float(amt_nt),
                        unit_price=float(p_nt) if p_nt is not None else None,
                        custom_emission_factor=float(gwp_nt) if gwp_nt is not None else None
                    ))
                except (ValueError, TypeError):
                    pass

            src_bm = ws_en.cell(r, 13).value
            amt_bm = ws_en.cell(r, 14).value
            p_bm = ws_en.cell(r, 15).value
            gwp_bm = ws_en.cell(r, 16).value
            if src_bm and amt_bm is not None:
                try:
                    en_bm.append(EnergyItem(
                        energy_source=str(src_bm).strip(),
                        amount_per_ref_volume=float(amt_bm),
                        unit_price=float(p_bm) if p_bm is not None else None,
                        custom_emission_factor=float(gwp_bm) if gwp_bm is not None else None
                    ))
                except (ValueError, TypeError):
                    pass

        sub_nt = SubsystemConfig(
            annual_facility_production=fac_prod_nt,
            capex_items=capex_nt,
            labor=labor_nt,
            materials=mat_nt,
            energy_items=en_nt,
            other_recurring_opex=other_rec_nt
        )
        sub_bm = SubsystemConfig(
            annual_facility_production=fac_prod_bm,
            capex_items=capex_bm,
            labor=labor_bm,
            materials=mat_bm,
            energy_items=en_bm,
            other_recurring_opex=other_rec_bm
        )

        return FullEntryInput(
            project=proj,
            market=mkt,
            functional_unit=fu,
            grid_scenario=GridScenario.BASELINE_2022,
            new_technology=sub_nt,
            commercial_benchmark=sub_bm
        )
